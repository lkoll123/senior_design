
from pathlib import Path
import io
import json

import boto3
import fitz
from openpyxl import load_workbook

# Configuration
SOURCE_DIR = Path("./data")
BUCKET = "your-s3-bucket"
PREFIX = "documents"

s3 = boto3.client("s3")


def parse_xlsx(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        result = {}

        for sheet in workbook:
            rows = [
                [str(cell) if cell is not None else None for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            result[sheet.title] = rows

        return result
    finally:
        workbook.close()


def parse_pdf(path):
    with fitz.open(path) as doc:
        return {
            "pages": [
                {"page": i + 1, "text": page.get_text()}
                for i, page in enumerate(doc)
            ]
        }


def ingest_file(path):
    relative_path = path.relative_to(SOURCE_DIR).as_posix()
    key = f"{PREFIX}/{relative_path}"

    # Upload original file
    s3.upload_file(str(path), BUCKET, f"raw/{key}")

    # Parse content
    if path.suffix.lower() == ".xlsx":
        content = parse_xlsx(path)
    elif path.suffix.lower() == ".pdf":
        content = parse_pdf(path)
    else:
        return

    # Upload extracted content as JSON
    payload = {
        "filename": path.name,
        "directory": str(path.relative_to(SOURCE_DIR).parent),
        "source_key": f"raw/{key}",
        "content": content,
    }

    json_key = f"parsed/{Path(key).with_suffix('.json').as_posix()}"

    s3.put_object(
        Bucket=BUCKET,
        Key=json_key,
        Body=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        ContentType="application/json",
    )

    print(f"Uploaded: {relative_path}")


def main():
    for path in SOURCE_DIR.rglob("*"):
        if path.is_file() and path.suffix.lower() in {".xlsx", ".pdf"}:
            try:
                ingest_file(path)
            except Exception as e:
                print(f"Failed: {path} — {e}")


if __name__ == "__main__":
    main()
